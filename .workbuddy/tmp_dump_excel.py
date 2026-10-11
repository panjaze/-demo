# -*- coding: utf-8 -*-
"""提取策划案 Excel 的全部文本内容（含合并单元格还原），供结构分析用"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

import openpyxl, xlrd

def dump_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        print(f"\n########## SHEET: {ws.title}  (rows={ws.max_row}, cols={ws.max_column}) ##########")
        merged = {}
        for m in ws.merged_cells.ranges:
            v = ws.cell(m.min_row, m.min_col).value
            for r in range(m.min_row, m.max_row + 1):
                for c in range(m.min_col, m.max_col + 1):
                    merged[(r, c)] = v
        for row in ws.iter_rows():
            cells = []
            for cell in row:
                v = cell.value
                if v is None:
                    v = merged.get((cell.row, cell.column))
                if v is None:
                    continue
                s = str(v).replace('\r', '')
                if len(s) > 400: s = s[:400] + f'…[len={len(s)}]'
                cells.append(f"[{cell.coordinate}] {s}")
            if cells:
                print(' | '.join(cells))

def dump_xls(path):
    wb = xlrd.open_workbook(path, formatting_info=False)
    for ws in wb.sheets():
        print(f"\n########## SHEET: {ws.name}  (rows={ws.nrows}, cols={ws.ncols}) ##########")
        for r in range(ws.nrows):
            cells = []
            for c in range(ws.ncols):
                v = ws.cell_value(r, c)
                if v in ('', None): continue
                s = str(v).replace('\r', '')
                if len(s) > 400: s = s[:400] + f'…[len={len(s)}]'
                col = xlrd.colname(c)
                cells.append(f"[{col}{r+1}] {s}")
            if cells:
                print(' | '.join(cells))

path = sys.argv[1]
if path.lower().endswith('.xls'):
    dump_xls(path)
else:
    dump_xlsx(path)
