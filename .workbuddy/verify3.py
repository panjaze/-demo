# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl
for f in ['游戏设计草案-拓边令.xlsx', '系统策划案-战斗与占领.xlsx', '系统策划案-政策·事件·任务与遗物.xlsx']:
    wb = openpyxl.load_workbook(f)
    info = []
    for name in wb.sheetnames:
        ws = wb[name]
        hs = [ws.row_dimensions[r].height for r in range(1, min(ws.max_row, 60)) if ws.row_dimensions[r].height]
        info.append(f"{name[:10]}({ws.max_row}行,行高{min(hs) if hs else '-'}~{max(hs) if hs else '-'})")
    print(f, '=>', ' | '.join(info))
